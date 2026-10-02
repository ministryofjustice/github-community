"""Read the old visibility audit spreadsheets and turn them into snapshots and events.

Pure functions only: no database or network access, so the import job can check
everything before it writes anything.

The spreadsheets have no GitHub ids, so repositories are matched by full name
("org/name"), case-insensitively. A rename is therefore a delete plus a create, and a
change that flips back between two files is lost. Both were agreed for this one-off import.
"""

import hashlib
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.projects.repository_stats.services.github_inventory import (
    RepositoryRecord,
    parse_github_datetime,
)
from app.projects.repository_stats.services.uk_time import london_date
from app.projects.repository_stats.services.visibility_logic import VisibilityEvent

IMPORT_SOURCE = "import"
BASELINE_FILE_NAME = "list_repos_17aug2026.xlsx"
BASELINE_DATE = date(2026, 8, 17)
VISIBILITIES = ("public", "private", "internal")

_DATED_FILE = re.compile(
    r"^(?P<kind>historical|visibility)-(?P<date>\d{4}-\d{2}-\d{2})\.xlsx$",
    re.IGNORECASE,
)


class AuditImportError(Exception):
    """The files can't be imported as they are. Nothing is written."""


@dataclass(frozen=True)
class AuditRow:
    org: str
    name: str
    visibility: str
    archived: bool | None = None
    fork: bool | None = None
    created_at: datetime | None = None
    pushed_at: datetime | None = None

    @property
    def key(self) -> str:
        return full_name_key(self.org, self.name)


@dataclass(frozen=True)
class AuditFile:
    captured_on: date
    file_name: str
    rows: tuple[AuditRow, ...]


@dataclass
class ImportPlan:
    """The snapshots and events for one organisation."""

    org: str
    snapshots: dict[date, list[RepositoryRecord]]
    events: list[VisibilityEvent]
    unresolved_ids: int = 0


def full_name_key(org: str, name: str) -> str:
    return f"{org}/{name}".lower()


def synthetic_github_id(org: str, name: str) -> int:
    """A stable stand-in id for a repository GitHub no longer has under this name.

    Negative, so it can never match a real GitHub id (always positive), and the same on
    every run, so re-running the import doesn't create new repositories. The scan job
    never sees these repositories, so it records them as deleted, which they are.
    """
    digest = hashlib.sha256(full_name_key(org, name).encode()).hexdigest()
    return -(int(digest[:15], 16) + 1)


# Reading the spreadsheets


def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def _header_index(headers: Sequence) -> dict[str, int]:
    return {
        _cell_text(header).lower(): index
        for index, header in enumerate(headers)
        if _cell_text(header)
    }


def _to_bool(value, file_name: str) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = _cell_text(value).lower()
    if text in ("true", "yes", "1"):
        return True
    if text in ("false", "no", "0"):
        return False
    raise AuditImportError(f"{file_name}: unexpected true/false value")


def _to_datetime(value) -> datetime | None:
    if value is None or value == "" or _cell_text(value).upper() == "N/A":
        return None
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.astimezone(UTC).replace(tzinfo=None)
        return value
    try:
        return parse_github_datetime(_cell_text(value))
    except ValueError:
        return None


def _visibility(value, file_name: str) -> str:
    visibility = _cell_text(value).lower()
    if visibility not in VISIBILITIES:
        raise AuditImportError(f"{file_name}: unexpected visibility value")
    return visibility


def _sheet(workbook, wanted: str, file_name: str):
    for sheet in workbook.worksheets:
        if sheet.title.strip().lower() == wanted.lower():
            return sheet
    raise AuditImportError(f"{file_name}: no '{wanted}' sheet")


def _rows(sheet) -> tuple[dict[str, int], Iterable[tuple]]:
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if headers is None:
        raise AuditImportError(f"{sheet.title}: empty sheet")
    return _header_index(headers), rows


def _require(columns: Mapping[str, int], names: Iterable[str], file_name: str) -> None:
    missing = [name for name in names if name not in columns]
    if missing:
        raise AuditImportError(f"{file_name}: missing column(s) {', '.join(missing)}")


def read_baseline(path: Path, captured_on: date = BASELINE_DATE) -> AuditFile:
    """The full repository list ("Repos" sheet): org, repo, visibility, archived, fork,
    created_at and last_pushed_at. Every other column is ignored."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        columns, rows = _rows(_sheet(workbook, "Repos", path.name))
        _require(columns, ("org", "repo", "visibility"), path.name)

        def get(row, column):
            index = columns.get(column)
            return row[index] if index is not None and index < len(row) else None

        result = []
        for row in rows:
            org, name = _cell_text(get(row, "org")), _cell_text(get(row, "repo"))
            if not org and not name:
                continue
            if not org or not name:
                raise AuditImportError(f"{path.name}: a row has no org or repo")
            result.append(
                AuditRow(
                    org=org,
                    name=name,
                    visibility=_visibility(get(row, "visibility"), path.name),
                    archived=_to_bool(get(row, "archived"), path.name),
                    fork=_to_bool(get(row, "fork"), path.name),
                    created_at=_to_datetime(get(row, "created_at")),
                    pushed_at=_to_datetime(
                        get(row, "last_pushed_at") or get(row, "pushed_at")
                    ),
                )
            )
    finally:
        workbook.close()
    return AuditFile(captured_on, path.name, tuple(result))


def _current_visibility_column(columns: Mapping[str, int], captured_on: date) -> int:
    """'visibility_current_2026-09-18' in the daily files, 'visibility_aug24' in the
    historical one. A dated header must match the file name's date."""
    for header, index in columns.items():
        if header.startswith("visibility_current"):
            stated = header.removeprefix("visibility_current").strip("_ ")
            if stated and stated[:10] != captured_on.isoformat():
                raise AuditImportError(
                    f"'{header}' doesn't match the file date {captured_on.isoformat()}"
                )
            return index
    short = f"visibility_{captured_on:%b}{captured_on.day}".lower()
    if short in columns:
        return columns[short]
    raise AuditImportError(
        f"no current visibility column for {captured_on.isoformat()}"
    )


def read_comparison(path: Path, captured_on: date) -> AuditFile:
    """A full snapshot from the "Full Comparison" sheet: repo ("org/name") and the current
    visibility, plus archived when the file has it. Baseline columns are ignored."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        columns, rows = _rows(_sheet(workbook, "Full Comparison", path.name))
        _require(columns, ("repo",), path.name)
        try:
            visibility_index = _current_visibility_column(columns, captured_on)
        except AuditImportError as error:
            raise AuditImportError(f"{path.name}: {error}") from None
        repo_index = columns["repo"]
        archived_index = columns.get("archived")

        result = []
        for row in rows:
            full_name = _cell_text(row[repo_index] if repo_index < len(row) else None)
            if not full_name:
                continue
            org, separator, name = full_name.partition("/")
            if not separator or not org or not name or "/" in name:
                raise AuditImportError(f"{path.name}: a repo isn't 'org/name'")
            archived = None
            if archived_index is not None and archived_index < len(row):
                archived = _to_bool(row[archived_index], path.name)
            result.append(
                AuditRow(
                    org=org,
                    name=name,
                    visibility=_visibility(row[visibility_index], path.name),
                    archived=archived,
                )
            )
    finally:
        workbook.close()
    return AuditFile(captured_on, path.name, tuple(result))


def read_directory(
    directory: Path,
    baseline_file_name: str = BASELINE_FILE_NAME,
    baseline_date: date = BASELINE_DATE,
) -> tuple[list[AuditFile], list[str]]:
    """The baseline plus every historical-YYYY-MM-DD.xlsx and Visibility-YYYY-MM-DD.xlsx,
    in date order. Returns the files and the names of the ones it ignored."""
    if not directory.is_dir():
        raise AuditImportError(f"{directory} is not a directory")
    baseline_path = directory / baseline_file_name
    if not baseline_path.is_file():
        raise AuditImportError(f"No baseline file {baseline_file_name}")

    files = [read_baseline(baseline_path, baseline_date)]
    seen = {baseline_date: baseline_file_name}
    ignored = []
    for path in sorted(directory.iterdir()):
        if path.name == baseline_file_name or not path.is_file():
            continue
        match = _DATED_FILE.match(path.name)
        if not match:
            ignored.append(path.name)
            continue
        try:
            captured_on = date.fromisoformat(match["date"])
        except ValueError:
            raise AuditImportError(f"{path.name}: not a real date") from None
        if captured_on <= baseline_date:
            raise AuditImportError(f"{path.name}: not after the baseline date")
        if captured_on in seen:
            raise AuditImportError(
                f"{path.name} and {seen[captured_on]} are for the same date"
            )
        seen[captured_on] = path.name
        files.append(read_comparison(path, captured_on))
    return sorted(files, key=lambda file: file.captured_on), ignored


# Building snapshots and events


def _created_on(
    created_at: datetime | None, previous_on: date, occurred_on: date
) -> date:
    """The day GitHub says the repository was created, if that's between the two files;
    otherwise the date of the file it first appears in."""
    if created_at is not None:
        created_on = london_date(created_at)
        if previous_on < created_on <= occurred_on:
            return created_on
    return occurred_on


def _not_after(value: datetime | None, captured_on: date) -> datetime | None:
    """GitHub's current pushed_at can be later than an older snapshot; drop it then."""
    if value is not None and london_date(value) > captured_on:
        return None
    return value


def diff_records(
    before: Mapping[int, RepositoryRecord],
    after: Mapping[int, RepositoryRecord],
    previous_on: date,
    occurred_on: date,
) -> list[VisibilityEvent]:
    """Changed, created, archived and deleted events between two snapshots, by github_id.
    Changes are dated to the later snapshot; created events to the real creation day when
    GitHub knows it."""

    def event(record, event_type, on=occurred_on, from_vis=None, to_vis=None):
        return VisibilityEvent(
            github_id=record.github_id,
            org=record.org,
            name=record.name,
            event_type=event_type,
            occurred_on=on,
            source=IMPORT_SOURCE,
            from_visibility=from_vis,
            to_visibility=to_vis,
        )

    events = []
    for github_id, record in sorted(after.items(), key=lambda i: i[1].name.lower()):
        old = before.get(github_id)
        if old is None:
            on = _created_on(record.created_at, previous_on, occurred_on)
            events.append(event(record, "created", on, to_vis=record.visibility))
            continue
        if old.visibility != record.visibility:
            events.append(
                event(
                    record, "changed", from_vis=old.visibility, to_vis=record.visibility
                )
            )
        if record.archived and not old.archived:
            events.append(event(record, "archived"))
    for github_id, old in sorted(before.items(), key=lambda i: i[1].name.lower()):
        if github_id not in after:
            events.append(event(old, "deleted", from_vis=old.visibility))
    return sorted(events, key=lambda e: e.occurred_on)


def build_plan(
    org: str,
    files: Sequence[AuditFile],
    inventory: Mapping[str, RepositoryRecord],
) -> ImportPlan:
    """Snapshots for every file and the events between consecutive files, for one org.

    inventory is GitHub's current repository list keyed by full_name_key. It gives each
    repository its real github_id (so the scan job matches it later) and, for repositories
    created after the baseline, the real creation date. Values a file doesn't have
    (archived, fork, created_at, pushed_at) are carried forward from the previous file.
    The first file is the baseline: it has no events.
    """
    snapshots: dict[date, list[RepositoryRecord]] = {}
    events: list[VisibilityEvent] = []
    unresolved: set[int] = set()
    previous: dict[str, RepositoryRecord] = {}
    previous_on: date | None = None

    for file in files:
        current: dict[str, RepositoryRecord] = {}
        for row in file.rows:
            if row.org != org:
                continue
            if row.key in current:
                raise AuditImportError(
                    f"{file.file_name}: a repository is listed twice"
                )
            known = inventory.get(row.key)
            old = previous.get(row.key)
            github_id = (
                known.github_id if known else synthetic_github_id(row.org, row.name)
            )
            if known is None:
                unresolved.add(github_id)

            def pick(value, field, default=None, old=old, known=known):
                if value is not None:
                    return value
                if old is not None:
                    return getattr(old, field)
                if known is not None and field != "archived":
                    return getattr(known, field)
                return default

            current[row.key] = RepositoryRecord(
                github_id=github_id,
                org=row.org,
                name=row.name,
                visibility=row.visibility,
                archived=bool(pick(row.archived, "archived", False)),
                fork=bool(pick(row.fork, "fork", False)),
                created_at=pick(row.created_at, "created_at"),
                pushed_at=_not_after(
                    pick(row.pushed_at, "pushed_at"), file.captured_on
                ),
            )
        by_id = {record.github_id: record for record in current.values()}
        if len(by_id) != len(current):
            raise AuditImportError(f"{file.file_name}: two repositories share an id")
        snapshots[file.captured_on] = list(current.values())
        if previous_on is not None:
            events += diff_records(
                {r.github_id: r for r in previous.values()},
                by_id,
                previous_on,
                file.captured_on,
            )
        previous, previous_on = current, file.captured_on

    return ImportPlan(org, snapshots, events, len(unresolved))


def event_counts(events: Iterable[VisibilityEvent]) -> dict[str, int]:
    counts = {"changed": 0, "created": 0, "deleted": 0, "archived": 0}
    for event in events:
        counts[event.event_type] = counts.get(event.event_type, 0) + 1
    return counts
